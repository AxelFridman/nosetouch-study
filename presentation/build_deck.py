#!/usr/bin/env python3
"""Build the NoseTouch deck (.pptx) from presentation/OUTLINE.md with python-pptx.

This is the stub of the builder planned in docs/09 (D32). It already renders every slide of the
outline: the title, message bullets, visuals, a QR code for the predictor and the speaker notes. A
visual that does not exist yet becomes a labelled placeholder box, so the deck always builds.

Visuals are resolved in this order:
  1. a registry ID (F13, T09, ...): the slide-style render of the report's asset step,
     report/figures/slides/<ID>_*.png, then report/figures/<ID>_*.png (PNG, GIF or JPEG; PDF
     figures cannot go into PowerPoint);
  2. the outline's stand-in after '>', e.g. "F13 > analysis/reports/baseline_synthetic/roc.png";
  3. a placeholder box naming the ID.
A plain path (from the repository root) is used as is. {key} fields in titles and bullets are filled
from report/tables/numbers.json, which the report's asset step writes.

Usage, from the repository root (python-pptx is not in analysis/requirements.txt; see README.md):
    analysis/.venv/Scripts/python presentation/build_deck.py            # -> presentation/build/NoseTouch_deck.pptx
    analysis/.venv/Scripts/python presentation/build_deck.py --check    # parse and check only, write nothing
    analysis/.venv/Scripts/python presentation/build_deck.py --final    # stand-ins, placeholders, missing numbers = errors

Exit codes: 0 ok; 1 errors (draft: structural problems only; --final: every problem); 2 usage error or
python-pptx missing.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
DEFAULT_OUTLINE = HERE / 'OUTLINE.md'
DEFAULT_OUT = HERE / 'build' / 'NoseTouch_deck.pptx'
REPORT_TOML = REPO / 'report' / 'report.toml'

INK = (0x1A, 0x1A, 0x1A)
MUTED = (0x6B, 0x6B, 0x6B)
ACCENT = (0x00, 0x79, 0x8C)            # RIGHT hand colour, used as the deck accent
PLACEHOLDER_FILL = (0xF2, 0xF2, 0xF2)
PLACEHOLDER_LINE = (0xA0, 0xA0, 0xA0)
NOTE_MAX = 160
MAX_DECK_BYTES = 50 * 1024 * 1024
SLIDE_W, SLIDE_H = 13.333, 7.5          # inches, 16:9
IMAGE_EXT = ('.png', '.gif', '.jpg', '.jpeg')
LAYOUTS = ('title', 'figure', 'two', 'demo', 'text')

SLIDE_RE = re.compile(r'^###\s+([SB]\d{2})\s*[·|]\s*(.+?)\s*$')
FIELD_RE = re.compile(r'^-\s+([a-z_]+):\s*(.*)$')
KEY_RE = re.compile(r'\{([A-Za-z0-9_]+(?:/[A-Za-z0-9_]+)*)\}')
ID_RE = re.compile(r'^[FT][0-9A-Z]{1,3}$')
USE_IF_RE = re.compile(r'^([A-Za-z0-9_/]+)\s*(<=|>=|<|>)\s*([-+]?[0-9.]+)$')


@dataclass
class Slide:
    sid: str
    title: str
    fields: dict[str, list[str]] = field(default_factory=dict)

    def first(self, key: str, default: str = '') -> str:
        values = self.fields.get(key)
        return values[0] if values else default

    def every(self, key: str) -> list[str]:
        return self.fields.get(key, [])


@dataclass
class Visual:
    token: str
    label: str
    path: Path | None
    standin: bool = False


@dataclass
class Problem:
    sid: str
    severity: str  # 'error' always fails; 'warning' fails only with --final
    text: str


# ───────────────────────────── outline and numbers ─────────────────────────────


def parse_outline(path: Path) -> tuple[dict[str, str], list[Slide]]:
    """Settings (the '## Settings' list) and slides ('### S01 · title' plus '- key: value' lines)."""
    settings: dict[str, str] = {}
    slides: list[Slide] = []
    section = ''
    current: Slide | None = None
    for raw in path.read_text(encoding='utf-8').splitlines():
        line = raw.rstrip()
        if line.startswith('## '):
            section, current = line[3:].strip().lower(), None
            continue
        m = SLIDE_RE.match(line)
        if m:
            current = Slide(m.group(1), m.group(2))
            slides.append(current)
            continue
        f = FIELD_RE.match(line)
        if not f:
            continue
        key, value = f.group(1), f.group(2).strip()
        if current is not None:
            current.fields.setdefault(key, []).append(value)
        elif section == 'settings':
            settings[key] = value
    return settings, slides


def load_numbers(path: Path) -> dict[str, str]:
    if not path.is_file():
        return {}
    data = json.loads(path.read_text(encoding='utf-8'))
    return {str(k): str(v) for k, v in data.items()} if isinstance(data, dict) else {}


def predictor_url(settings: dict[str, str]) -> str:
    """report/report.toml [predictor].url (one link for report, PI sheet and deck), else the outline's."""
    try:
        import tomllib  # Python 3.11+
        with REPORT_TOML.open('rb') as fh:
            url = tomllib.load(fh).get('predictor', {}).get('url')
        if url:
            return str(url)
    except (ImportError, OSError, ValueError):
        pass
    return settings.get('predictor_url', '')


def fill(text: str, numbers: dict[str, str], missing: set[str]) -> str:
    def repl(m: re.Match[str]) -> str:
        if m.group(1) in numbers:
            return numbers[m.group(1)]
        missing.add(m.group(1))
        return m.group(0)
    return KEY_RE.sub(repl, text)


def condition(expr: str, numbers: dict[str, str]) -> bool | None:
    """'calib/ece <= 0.05' -> True/False; None when the key or the value is missing."""
    m = USE_IF_RE.match(expr.strip())
    if not m or m.group(1) not in numbers:
        return None
    try:
        value, limit = float(numbers[m.group(1)]), float(m.group(3))
    except ValueError:
        return None
    return {'<=': value <= limit, '>=': value >= limit, '<': value < limit, '>': value > limit}[m.group(2)]


def slide_title(s: Slide, numbers: dict[str, str], problems: list[Problem]) -> str:
    missing: set[str] = set()
    title = fill(s.title, numbers, missing)
    use_if = s.first('use_if')
    ok = condition(use_if, numbers) if use_if else True
    if (missing or ok is not True) and s.first('fallback'):
        why = f'missing {sorted(missing)}' if missing else f'condition "{use_if}" not met or unknown'
        problems.append(Problem(s.sid, 'warning', f'title uses the fallback ({why})'))
        return s.first('fallback')
    if missing:
        problems.append(Problem(s.sid, 'warning', f'unresolved numbers in the title: {sorted(missing)}'))
    return title


def resolve_visual(token: str, figure_dirs: list[Path]) -> Visual:
    primary, _, standin = (part.strip() for part in token.partition('>'))
    if ID_RE.match(primary):
        for d in figure_dirs:
            for ext in IMAGE_EXT:
                hits = sorted(d.glob(f'{primary}_*{ext}')) + sorted(d.glob(f'{primary}{ext}'))
                if hits:
                    return Visual(token, primary, hits[0])
        if standin:
            p = REPO / standin
            if p.is_file() and p.suffix.lower() in IMAGE_EXT:
                return Visual(token, f'{primary} (stand-in)', p, standin=True)
        return Visual(token, primary, None)
    p = REPO / primary
    return Visual(token, primary, p if p.is_file() and p.suffix.lower() in IMAGE_EXT else None)


def check(slides: list[Slide], numbers: dict[str, str], figure_dirs: list[Path]) -> list[Problem]:
    problems: list[Problem] = []
    seen: set[str] = set()
    for s in slides:
        if s.sid in seen:
            problems.append(Problem(s.sid, 'error', 'duplicate slide id'))
        seen.add(s.sid)
        if not s.title.strip():
            problems.append(Problem(s.sid, 'error', 'no title'))
        slide_title(s, numbers, problems)  # records fallback and missing-number warnings
        layout = s.first('layout', 'figure')
        if layout not in LAYOUTS:
            problems.append(Problem(s.sid, 'error', f'unknown layout {layout!r}'))
        notes = s.first('notes')
        if not notes:
            problems.append(Problem(s.sid, 'error', 'no speaker note'))
        elif len(notes) > NOTE_MAX:
            problems.append(Problem(s.sid, 'error', f'speaker note has {len(notes)} > {NOTE_MAX} characters'))
        missing: set[str] = set()
        for msg in s.every('message'):
            fill(msg, numbers, missing)
        if missing:
            problems.append(Problem(s.sid, 'warning', f'unresolved numbers in bullets: {sorted(missing)}'))
        for vis in s.every('visual'):
            for token in filter(None, (t.strip() for t in vis.split(';'))):
                v = resolve_visual(token, figure_dirs)
                if v.path is None:
                    problems.append(Problem(s.sid, 'warning', f'visual {v.label} not found (placeholder drawn)'))
                elif v.standin:
                    problems.append(Problem(s.sid, 'warning', f'visual {v.label}: stand-in {v.path.relative_to(REPO)}'))
    return problems


# ───────────────────────────── rendering ─────────────────────────────


def render(slides: list[Slide], settings: dict[str, str], numbers: dict[str, str], figure_dirs: list[Path],
           out: Path, problems: list[Problem]) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    try:
        from PIL import Image
    except ImportError:  # python-pptx needs Pillow, so this is only a safety net
        Image = None

    out.parent.mkdir(parents=True, exist_ok=True)
    url = predictor_url(settings)
    qr_png: Path | None = out.parent / 'qr_predictor.png'
    try:
        import segno
        if not url:
            raise ValueError('no predictor URL')
        segno.make(url, error='m').save(str(qr_png), scale=12, border=4, dark='#111111')
    except (ImportError, ValueError) as exc:
        qr_png = None
        problems.append(Problem('deck', 'warning', f'QR codes drawn as placeholders ({exc})'))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    def text_box(slide, x, y, w, h, lines, size, *, bold=False, color=INK, bullets=False, align=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if align is not None:
                p.alignment = align
            if bullets and not re.match(r'^\d+\.', line):
                line = '• ' + line
            run = p.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(*color)
            p.space_after = Pt(size * 0.5)
        return box

    def placeholder(slide, x, y, w, h, label):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*PLACEHOLDER_FILL)
        shape.line.color.rgb = RGBColor(*PLACEHOLDER_LINE)
        tf = shape.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        tf.text = f'{label}\nnot generated yet'
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(16)
                r.font.color.rgb = RGBColor(*MUTED)

    def picture(slide, path, label, x, y, w, h):
        """Fit the image into the box, keeping its aspect ratio; centred."""
        if path is None:
            placeholder(slide, x, y, w, h, label)
            return
        iw, ih = (1, 1)
        if Image is not None:
            with Image.open(path) as im:
                iw, ih = im.size
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                                 Inches(pw), Inches(ph))

    def visuals_of(s):
        out_list = []
        for vis in s.every('visual'):
            for token in filter(None, (t.strip() for t in vis.split(';'))):
                v = resolve_visual(token, figure_dirs)
                out_list.append(v)
        return out_list

    def qr(slide, x, y, size):
        if qr_png is not None and qr_png.is_file():
            slide.shapes.add_picture(str(qr_png), Inches(x), Inches(y), Inches(size), Inches(size))
        else:
            placeholder(slide, x, y, size, size, 'QR code')

    footer_template = settings.get('footer', '')
    for s in slides:
        slide = prs.slides.add_slide(blank)
        layout = s.first('layout', 'figure')
        title = slide_title(s, numbers, [])  # its problems were recorded by check()
        missing: set[str] = set()
        messages = [fill(m, numbers, missing) for m in s.every('message')]
        vis = visuals_of(s)

        if layout == 'title':
            text_box(slide, 0.6, 0.7, 12.1, 1.5, [title], 40, bold=True)
            text_box(slide, 0.6, 2.2, 12.1, 1.2, messages, 20, color=MUTED)
            slots = [(0.6 + i * 4.3, 3.5, 4.0, 3.4) for i in range(min(len(vis), 2))]
            for v, (x, y, w, h) in zip(vis, slots):
                picture(slide, v.path, v.label, x, y, w, h)
            if s.first('qr'):
                qr(slide, 10.2, 4.0, 2.6)
                text_box(slide, 9.6, 6.7, 3.6, 0.4, [url], 11, color=MUTED, align=PP_ALIGN.CENTER)
        elif layout == 'demo':
            text_box(slide, 0.6, 0.35, 12.1, 1.1, [title], 30, bold=True)
            qr(slide, 0.6, 1.5, 4.6)
            text_box(slide, 0.6, 6.25, 4.6, 0.6, [url], 12, color=MUTED, align=PP_ALIGN.CENTER)
            text_box(slide, 5.6, 1.6, 4.2, 5.0, messages, 18, bullets=True)
            if vis:
                picture(slide, vis[0].path, vis[0].label, 10.1, 1.5, 2.7, 5.4)
        elif layout == 'two':
            text_box(slide, 0.6, 0.35, 12.1, 1.1, [title], 30, bold=True)
            slots = [(0.6, 1.6, 5.9, 4.1), (6.83, 1.6, 5.9, 4.1)]
            for v, (x, y, w, h) in zip(vis, slots):
                picture(slide, v.path, v.label, x, y, w, h)
            text_box(slide, 0.6, 5.85, 12.1, 1.1, messages, 18, bullets=True)
        elif layout == 'text':
            text_box(slide, 0.6, 0.35, 12.1, 1.1, [title], 30, bold=True)
            text_box(slide, 0.9, 1.8, 11.5, 4.8, messages, 22, bullets=True)
        else:  # figure: one visual on the left, up to 3 bullets on the right
            text_box(slide, 0.6, 0.35, 12.1, 1.1, [title], 30, bold=True)
            if vis:
                picture(slide, vis[0].path, vis[0].label, 0.6, 1.6, 7.9, 5.2)
            text_box(slide, 8.8, 1.8, 3.95, 4.8, messages, 18, bullets=True)

        if s.first('footer') == 'yes' and footer_template:
            text_box(slide, 0.6, 7.0, 9.5, 0.35, [fill(footer_template, numbers, missing)], 10, color=MUTED)
        text_box(slide, 12.0, 7.0, 0.8, 0.35, [f'{s.sid}'], 10, color=MUTED, align=PP_ALIGN.RIGHT)
        slide.notes_slide.notes_text_frame.text = s.first('notes')
        if s.sid == 'S01':  # thin accent bar under the title
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(2.1), Inches(1.6), Inches(0.06))
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGBColor(*ACCENT)
            bar.line.fill.background()

    prs.save(str(out))
    if out.stat().st_size > MAX_DECK_BYTES:
        problems.append(Problem('deck', 'warning', f'{out.name} is larger than 50 MB'))


# ───────────────────────────── command line ─────────────────────────────


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(errors='replace')  # Windows consoles: never crash on a non-ASCII title
    ap = argparse.ArgumentParser(description='Build the NoseTouch deck from presentation/OUTLINE.md.')
    ap.add_argument('--outline', type=Path, default=DEFAULT_OUTLINE, help='outline file (default: %(default)s)')
    ap.add_argument('--numbers', type=Path, help='numbers.json (default: the outline setting "numbers")')
    ap.add_argument('--out', type=Path, default=DEFAULT_OUT, help='output .pptx (default: %(default)s)')
    ap.add_argument('--final', action='store_true', help='treat every problem as an error')
    ap.add_argument('--check', action='store_true', help='parse and check only; write nothing')
    ap.add_argument('--no-backup', action='store_true', help='leave out the backup (B) slides')
    args = ap.parse_args(argv)

    if not args.outline.is_file():
        print(f'error: outline not found: {args.outline}')
        return 2
    settings, slides = parse_outline(args.outline)
    if args.no_backup:
        slides = [s for s in slides if s.sid.startswith('S')]
    if not slides:
        print(f'error: no slides found in {args.outline}')
        return 2
    numbers_path = args.numbers or (REPO / settings.get('numbers', 'report/tables/numbers.json'))
    numbers = load_numbers(numbers_path)
    figure_dirs = [REPO / d.strip() for d in settings.get('figures', 'report/figures').split(';') if d.strip()]

    problems = check(slides, numbers, figure_dirs)
    if not args.check:
        try:
            import pptx  # noqa: F401
        except ImportError:
            print('error: python-pptx is not installed. Install it into the analysis venv:\n'
                  '  analysis\\.venv\\Scripts\\python -m pip install python-pptx==1.0.2 segno==1.6.6')
            return 2
        render(slides, settings, numbers, figure_dirs, args.out, problems)

    errors = [p for p in problems if p.severity == 'error' or args.final]
    for p in problems:
        flag = 'ERROR' if (p.severity == 'error' or args.final) else 'warn '
        print(f'{flag} {p.sid}: {p.text}')
    n_main = sum(s.sid.startswith('S') for s in slides)
    print(f'{len(slides)} slides ({n_main} main, {len(slides) - n_main} backup); '
          f'{len(numbers)} numbers from {numbers_path.relative_to(REPO) if numbers_path.is_relative_to(REPO) else numbers_path}; '
          f'{len(problems)} problem(s), {len(errors)} error(s)')
    if not args.check:
        print(f'wrote {args.out}')
    return 1 if errors else 0


if __name__ == '__main__':
    sys.exit(main())
